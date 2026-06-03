"""Generate CSV, Excel, JSON, and HTML reports for analyzed road videos."""

from __future__ import annotations

import html
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from .defect_triage import write_defect_report
from .test_case_generator import write_test_cases


EVENT_COLUMNS = [
    "video_name",
    "frame_number",
    "timestamp_sec",
    "event_type",
    "risk_level",
    "detected_objects",
    "object_count",
    "brightness",
    "evidence_image",
    "tester_note",
]


def build_summary(
    video_name: str,
    total_frames_analyzed: int,
    events: list[dict[str, Any]],
    detected_object_counts: dict[str, int] | None = None,
    analysis_interval_sec: float = 1.0,
    model: str = "yolov8n.pt",
    confidence_threshold: float = 0.35,
) -> dict[str, Any]:
    risk_level_counts = Counter(event["risk_level"] for event in events)
    event_type_counts = Counter(event["event_type"] for event in events)
    return {
        "video_name": video_name,
        "analyzed_frames": total_frames_analyzed,
        "total_frames_analyzed": total_frames_analyzed,
        "total_events": len(events),
        "total_risk_events": len(events),
        "high_count": risk_level_counts.get("High", 0),
        "medium_count": risk_level_counts.get("Medium", 0),
        "low_count": risk_level_counts.get("Low", 0),
        "risk_level_counts": {
            "High": risk_level_counts.get("High", 0),
            "Medium": risk_level_counts.get("Medium", 0),
            "Low": risk_level_counts.get("Low", 0),
        },
        "event_type_counts": dict(event_type_counts),
        "detected_object_counts": detected_object_counts or {},
        "analysis_interval_sec": analysis_interval_sec,
        "model": model,
        "confidence_threshold": confidence_threshold,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


def write_reports(events: list[dict[str, Any]], summary: dict[str, Any], output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "event_log.csv"
    excel_path = output_dir / "risk_events.xlsx"
    summary_path = output_dir / "summary.json"
    html_path = output_dir / "test_report.html"

    frame = pd.DataFrame(events, columns=EVENT_COLUMNS)
    frame.to_csv(csv_path, index=False)
    frame.to_excel(excel_path, index=False)
    _adjust_excel_columns(excel_path)

    summary["output_files"] = {
        "event_log_csv": str(csv_path).replace("\\", "/"),
        "risk_events_xlsx": str(excel_path).replace("\\", "/"),
        "summary_json": str(summary_path).replace("\\", "/"),
        "test_report_html": str(html_path).replace("\\", "/"),
    }

    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    html_path.write_text(_render_html(frame, summary), encoding="utf-8")

    return {
        "csv": csv_path,
        "excel": excel_path,
        "summary": summary_path,
        "html": html_path,
    }


def _adjust_excel_columns(excel_path: Path) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return

    workbook = load_workbook(excel_path)
    worksheet = workbook.active
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 48)
    workbook.save(excel_path)


def _render_count_items(counts: dict[str, int]) -> str:
    if not counts:
        return "<p>No events detected.</p>"
    return "".join(
        f"<div class=\"metric\"><span>{html.escape(str(name))}</span><strong>{count}</strong></div>"
        for name, count in sorted(counts.items())
    )


def _relative_evidence_path(evidence_image: str) -> str:
    if not evidence_image:
        return ""
    path = Path(evidence_image)
    if len(path.parts) >= 2 and path.parts[0] == "outputs":
        return str(Path(*path.parts[1:])).replace("\\", "/")
    return evidence_image.replace("\\", "/")


def _render_rows(frame: pd.DataFrame) -> str:
    if frame.empty:
        return "<tr><td colspan=\"10\">No risk events were detected.</td></tr>"

    rows = []
    for record in frame.to_dict(orient="records"):
        evidence = _relative_evidence_path(str(record.get("evidence_image") or ""))
        evidence_html = (
            f'<a href="{html.escape(evidence)}"><img src="{html.escape(evidence)}" alt="Evidence frame"></a>'
            if evidence
            else ""
        )
        rows.append(
            "<tr>"
            f"<td>{html.escape(str(record['frame_number']))}</td>"
            f"<td>{html.escape(str(record['timestamp_sec']))}</td>"
            f"<td><span class=\"badge {html.escape(str(record['risk_level']).lower())}\">{html.escape(str(record['risk_level']))}</span></td>"
            f"<td>{html.escape(str(record['event_type']))}</td>"
            f"<td>{html.escape(str(record['detected_objects']))}</td>"
            f"<td>{html.escape(str(record['brightness']))}</td>"
            f"<td>{evidence_html}</td>"
            f"<td>{html.escape(str(record['tester_note']))}</td>"
            "</tr>"
        )
    return "\n".join(rows)


def _render_html(frame: pd.DataFrame, summary: dict[str, Any]) -> str:
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>FSD Road Scene Test Logger Report</title>
  <style>
    :root {{
      --ink: #172026;
      --muted: #66727d;
      --line: #d9e0e6;
      --panel: #f6f8fa;
      --high: #b42318;
      --medium: #b54708;
      --low: #175cd3;
    }}
    body {{
      margin: 0;
      color: var(--ink);
      background: #ffffff;
      font-family: "Segoe UI", Tahoma, sans-serif;
      line-height: 1.5;
    }}
    main {{
      max-width: 1180px;
      margin: 0 auto;
      padding: 32px 20px 56px;
    }}
    header {{
      border-bottom: 1px solid var(--line);
      margin-bottom: 24px;
      padding-bottom: 18px;
    }}
    h1, h2 {{
      margin: 0 0 10px;
      letter-spacing: 0;
    }}
    .meta {{
      color: var(--muted);
      margin: 0;
    }}
    .summary {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 12px;
      margin: 22px 0;
    }}
    .metric {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 14px;
    }}
    .metric span {{
      color: var(--muted);
      display: block;
      font-size: 0.9rem;
    }}
    .metric strong {{
      display: block;
      font-size: 1.55rem;
      margin-top: 4px;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin-top: 12px;
      font-size: 0.92rem;
    }}
    th, td {{
      border-bottom: 1px solid var(--line);
      padding: 10px;
      text-align: left;
      vertical-align: top;
    }}
    th {{
      background: var(--panel);
      font-weight: 650;
    }}
    img {{
      max-width: 180px;
      border: 1px solid var(--line);
      border-radius: 4px;
    }}
    .badge {{
      border-radius: 999px;
      color: white;
      display: inline-block;
      font-size: 0.78rem;
      font-weight: 700;
      padding: 3px 9px;
    }}
    .high {{ background: var(--high); }}
    .medium {{ background: var(--medium); }}
    .low {{ background: var(--low); }}
    section {{
      margin-top: 28px;
    }}
  </style>
</head>
<body>
  <main>
    <header>
      <h1>FSD Road Scene Test Logger Report</h1>
      <p class="meta">Video analyzed: {html.escape(str(summary["video_name"]))}</p>
      <p class="meta">Generated at: {html.escape(str(summary["generated_at"]))}</p>
    </header>

    <section class="summary">
      <div class="metric"><span>Total frames analyzed</span><strong>{summary["total_frames_analyzed"]}</strong></div>
      <div class="metric"><span>Total risk events</span><strong>{summary["total_risk_events"]}</strong></div>
      <div class="metric"><span>Interval seconds</span><strong>{summary["analysis_interval_sec"]}</strong></div>
      <div class="metric"><span>Confidence threshold</span><strong>{summary["confidence_threshold"]}</strong></div>
      {_render_count_items(summary["risk_level_counts"])}
    </section>

    <section>
      <h2>Analysis Parameters</h2>
      <div class="summary">
        <div class="metric"><span>Model</span><strong>{html.escape(str(summary["model"]))}</strong></div>
        <div class="metric"><span>Video</span><strong>{html.escape(str(summary["video_name"]))}</strong></div>
      </div>
    </section>

    <section>
      <h2>Event Type Counts</h2>
      <div class="summary">{_render_count_items(summary["event_type_counts"])}</div>
    </section>

    <section>
      <h2>Risk Event Log</h2>
      <table>
        <thead>
          <tr>
            <th>Frame</th>
            <th>Time (s)</th>
            <th>Risk</th>
            <th>Event</th>
            <th>Detected Objects</th>
            <th>Brightness</th>
            <th>Evidence</th>
            <th>Tester Note</th>
          </tr>
        </thead>
        <tbody>
          {_render_rows(frame)}
        </tbody>
      </table>
    </section>

    <section>
      <h2>Testing Interpretation</h2>
      <p>This report converts sampled road-video frames into structured ADAS-style risk observations. Each event includes a timestamp, detected objects, a rule-based risk level, and an evidence frame when available.</p>
    </section>

    <section>
      <h2>Limitations</h2>
      <p>This is a lightweight offline analysis tool. It uses a pretrained object detector and simplified rules, so results can include false positives or missed detections and should not be treated as a real autonomous driving decision system.</p>
    </section>
  </main>
</body>
</html>
"""


def write_platform_reports(
    events: list[dict[str, Any]],
    defects: list[dict[str, Any]],
    test_cases: list[dict[str, Any]],
    output_dir: Path,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Path]:
    """Write ADAS scenario mining platform outputs."""
    output_dir.mkdir(parents=True, exist_ok=True)
    charts_dir = output_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)

    event_log_path = output_dir / "event_log.csv"
    risk_events_path = output_dir / "risk_events.xlsx"
    summary_path = output_dir / "summary.json"
    html_path = output_dir / "test_report.html"

    event_frame = pd.DataFrame(events)
    defect_frame = pd.DataFrame(defects)
    case_frame = pd.DataFrame(test_cases)

    event_frame.to_csv(event_log_path, index=False)
    event_frame.to_excel(risk_events_path, index=False)
    _adjust_excel_columns(risk_events_path)

    test_case_paths = write_test_cases(test_cases, output_dir)
    defect_report_path = write_defect_report(defects, output_dir / "defect_report.md")

    chart_paths = _write_platform_charts(event_frame, defect_frame, charts_dir)
    summary = _build_platform_summary(events, defects, test_cases, metadata or {}, chart_paths)
    summary["output_files"] = {
        "event_log_csv": str(event_log_path).replace("\\", "/"),
        "risk_events_xlsx": str(risk_events_path).replace("\\", "/"),
        "test_cases_xlsx": str(test_case_paths["xlsx"]).replace("\\", "/"),
        "test_cases_csv": str(test_case_paths["csv"]).replace("\\", "/"),
        "defect_report_md": str(defect_report_path).replace("\\", "/"),
        "summary_json": str(summary_path).replace("\\", "/"),
        "test_report_html": str(html_path).replace("\\", "/"),
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    html_path.write_text(
        _render_platform_html(event_frame, defect_frame, case_frame, summary, chart_paths),
        encoding="utf-8",
    )
    return {
        "event_log": event_log_path,
        "risk_events": risk_events_path,
        "test_cases_xlsx": test_case_paths["xlsx"],
        "test_cases_csv": test_case_paths["csv"],
        "defect_report": defect_report_path,
        "summary": summary_path,
        "html": html_path,
    }


def _build_platform_summary(
    events: list[dict[str, Any]],
    defects: list[dict[str, Any]],
    test_cases: list[dict[str, Any]],
    metadata: dict[str, Any],
    chart_paths: dict[str, Path],
) -> dict[str, Any]:
    risk_counts = Counter(event.get("risk_level", "LOW") for event in events)
    scenario_counts = Counter(event.get("scenario_type", "unknown") for event in events)
    priority_counts = Counter(defect.get("priority", "P3") for defect in defects)
    return {
        "project_name": "ADAS Scenario-Based Risk Mining & Regression Test Platform",
        "metadata": metadata,
        "total_risk_events": len(events),
        "total_defects": len(defects),
        "total_test_cases": len(test_cases),
        "critical_events": risk_counts.get("CRITICAL", 0),
        "risk_level_counts": {
            "CRITICAL": risk_counts.get("CRITICAL", 0),
            "HIGH": risk_counts.get("HIGH", 0),
            "MEDIUM": risk_counts.get("MEDIUM", 0),
            "LOW": risk_counts.get("LOW", 0),
        },
        "scenario_type_counts": dict(scenario_counts),
        "priority_counts": dict(priority_counts),
        "chart_files": {name: str(path).replace("\\", "/") for name, path in chart_paths.items()},
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


def _write_platform_charts(event_frame: pd.DataFrame, defect_frame: pd.DataFrame, charts_dir: Path) -> dict[str, Path]:
    chart_paths: dict[str, Path] = {}
    try:
        import matplotlib.pyplot as plt
    except ImportError:
        return chart_paths

    if not event_frame.empty and "risk_level" in event_frame:
        path = charts_dir / "risk_level_distribution.png"
        event_frame["risk_level"].value_counts().plot(kind="bar", color=["#b42318", "#d97706", "#175cd3", "#667085"])
        plt.title("Risk Level Distribution")
        plt.tight_layout()
        plt.savefig(path)
        plt.close()
        chart_paths["risk_level_distribution"] = path

    if not event_frame.empty and "scenario_type" in event_frame:
        path = charts_dir / "scenario_type_distribution.png"
        event_frame["scenario_type"].value_counts().plot(kind="bar", color="#175cd3")
        plt.title("Scenario Type Distribution")
        plt.tight_layout()
        plt.savefig(path)
        plt.close()
        chart_paths["scenario_type_distribution"] = path

    if not defect_frame.empty and "priority" in defect_frame:
        path = charts_dir / "priority_distribution.png"
        defect_frame["priority"].value_counts().plot(kind="bar", color="#b54708")
        plt.title("Defect Priority Distribution")
        plt.tight_layout()
        plt.savefig(path)
        plt.close()
        chart_paths["priority_distribution"] = path

    return chart_paths


def _html_table(frame: pd.DataFrame, columns: list[str], empty_message: str) -> str:
    if frame.empty:
        return f"<p>{html.escape(empty_message)}</p>"
    visible = frame[[column for column in columns if column in frame.columns]].head(25)
    rows = []
    for record in visible.to_dict(orient="records"):
        cells = "".join(f"<td>{html.escape(str(value))}</td>" for value in record.values())
        rows.append(f"<tr>{cells}</tr>")
    headers = "".join(f"<th>{html.escape(column)}</th>" for column in visible.columns)
    return f"<table><thead><tr>{headers}</tr></thead><tbody>{''.join(rows)}</tbody></table>"


def _render_platform_html(
    event_frame: pd.DataFrame,
    defect_frame: pd.DataFrame,
    case_frame: pd.DataFrame,
    summary: dict[str, Any],
    chart_paths: dict[str, Path],
) -> str:
    chart_images = "".join(
        f'<img class="chart" src="{html.escape(str(path.relative_to(path.parents[1])).replace("\\\\", "/"))}" alt="{html.escape(name)}">'
        for name, path in chart_paths.items()
    )
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>ADAS Scenario-Based Risk Mining & Regression Test Report</title>
  <style>
    body {{ margin: 0; font-family: "Segoe UI", Tahoma, sans-serif; color: #172026; background: #fff; }}
    main {{ max-width: 1280px; margin: 0 auto; padding: 32px 20px 56px; }}
    header {{ border-bottom: 1px solid #d9e0e6; padding-bottom: 18px; }}
    h1, h2 {{ margin: 0 0 12px; }}
    .meta {{ color: #66727d; }}
    .grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 12px; margin: 22px 0; }}
    .metric {{ border: 1px solid #d9e0e6; border-radius: 6px; padding: 14px; background: #f6f8fa; }}
    .metric span {{ display: block; color: #66727d; font-size: 0.9rem; }}
    .metric strong {{ display: block; font-size: 1.6rem; }}
    table {{ border-collapse: collapse; width: 100%; font-size: 0.9rem; }}
    th, td {{ border-bottom: 1px solid #d9e0e6; padding: 9px; text-align: left; vertical-align: top; }}
    th {{ background: #f6f8fa; }}
    .chart {{ max-width: 420px; width: 100%; border: 1px solid #d9e0e6; border-radius: 6px; margin: 8px; }}
    section {{ margin-top: 28px; }}
  </style>
</head>
<body>
<main>
  <header>
    <h1>ADAS Scenario-Based Risk Mining & Regression Test Report</h1>
    <p class="meta">Generated at: {html.escape(summary["generated_at"])}</p>
  </header>
  <section class="grid">
    <div class="metric"><span>Total Risk Events</span><strong>{summary["total_risk_events"]}</strong></div>
    <div class="metric"><span>Critical Events</span><strong>{summary["critical_events"]}</strong></div>
    <div class="metric"><span>Generated Defects</span><strong>{summary["total_defects"]}</strong></div>
    <div class="metric"><span>Generated Test Cases</span><strong>{summary["total_test_cases"]}</strong></div>
  </section>
  <section><h2>Charts</h2>{chart_images or "<p>No charts generated.</p>"}</section>
  <section><h2>Top Risk Events</h2>{_html_table(event_frame.sort_values("risk_score", ascending=False) if "risk_score" in event_frame else event_frame, ["event_id", "scenario_type", "risk_level", "risk_score", "ttc", "thw", "description", "evidence_frame"], "No risk events.")}</section>
  <section><h2>Defect Triage Table</h2>{_html_table(defect_frame, ["defect_id", "linked_event_id", "priority", "feature", "scenario_type", "title"], "No defects.")}</section>
  <section><h2>Test Case Table</h2>{_html_table(case_frame, ["test_case_id", "scenario_type", "linked_event_id", "linked_defect_id", "priority", "expected_result"], "No test cases.")}</section>
  <section><h2>Limitations</h2><p>This first platform version uses deterministic synthetic detections and simplified scenario rules. It is intended for repeatable ADAS test workflow demonstration, not real vehicle safety validation.</p></section>
</main>
</body>
</html>
"""
